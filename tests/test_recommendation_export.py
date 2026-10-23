from copy import deepcopy
from datetime import datetime
from io import BytesIO
import unittest
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.models import RecommendationRequest
from app.services.preference_filters import RECOMMENDATION_VERSION
from app.services.recommender import RecommenderService


class RecommendationExportTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        self.client.cookies.set("vrs_auth", "1")
        self.payload = {
            "items": [{"institution": "示例大学", "major": "计算机科学与技术", "year": 2024,
                       "layer": "稳", "min_score": 0, "min_rank": 10000, "plan_count": None,
                       "subject_requirement_display": "物理和化学", "subject_eligibility": "unknown",
                       "subject_missing": ["化学"], "source_file": "专业分数线.xlsx", "source_sheet": "浙江",
                       "source_row": 123, "plan_source": "招生计划.xlsx", "plan_source_row": 42,
                       "fact_id": "offer-001", "plan_fact_id": "plan-001"}],
            "meta": {"reference_year": 2024, "recommendation_version": RECOMMENDATION_VERSION,
                     "input_snapshot": {"rank": 11100, "score": 0, "ratio_chong": 0},
                     "generated_at": "2026-10-23T10:00:00+00:00", "warnings": ["示例提示"]},
            "scope": "visible", "total_count": 3, "filters": {"keyword": "计算机", "layer": "稳"},
        }

    def export(self, payload=None):
        response = self.client.post("/api/v1/recommend/export", json=payload or self.payload)
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
        return response, load_workbook(BytesIO(response.content))

    def test_real_download_has_numeric_cells_sources_filter_and_frozen_header(self):
        response, workbook = self.export()
        self.addCleanup(workbook.close)
        self.assertIn("spreadsheetml.sheet", response.headers["content-type"])
        self.assertIn("attachment", response.headers["content-disposition"])
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertEqual(workbook.sheetnames, ["方案说明", "推荐清单"])
        sheet = workbook["推荐清单"]
        self.assertEqual(sheet["F2"].value, 0)
        self.assertEqual(sheet["G2"].value, 10000)
        self.assertEqual(sheet["G2"].data_type, "n")
        self.assertIsNone(sheet["H2"].value)
        self.assertEqual(sheet["N2"].value, "选科待核实")
        self.assertEqual(sheet["O2"].value, "化学")
        self.assertEqual([sheet.cell(2, col).value for col in range(16, 23)],
                         ["专业分数线.xlsx", "浙江", 123, "招生计划.xlsx", 42, "offer-001", "plan-001"])
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:V2")
        notes = dict(workbook["方案说明"].values)
        self.assertEqual(notes["导出条数 / 推荐总数"], "1 / 3")
        self.assertEqual(notes["生成时位次"], 11100)
        self.assertEqual(notes["生成时分数"], 0)
        self.assertEqual(notes["冲 / 稳 / 保比例"], "0 / 4 / 3")
        self.assertIn("计算机", notes["页面筛选"])
        self.assertEqual(notes["生成提示"], "示例提示")

    def test_formula_like_and_invalid_xml_text_are_literal(self):
        self.payload["items"][0].update(institution='=HYPERLINK("https://example.org", "x")',
                                        major="@SUM(1,2)\x00专业\uffff", source_file="+1+1", source_sheet="#REF!")
        self.payload["meta"]["warnings"] = ["=2+2"]
        _, workbook = self.export()
        self.addCleanup(workbook.close)
        for sheet in workbook:
            self.assertFalse(any(c.data_type in {"f", "e"} for row in sheet for c in row))
        self.assertTrue(workbook["推荐清单"]["C2"].value.startswith("=HYPERLINK"))
        self.assertEqual(workbook["推荐清单"]["D2"].value, "@SUM(1,2)专业")
        self.assertEqual(workbook["推荐清单"]["Q2"].value, "#REF!")

    def test_export_requires_login(self):
        self.client.cookies.clear()
        with patch("app.main.build_recommendation_workbook") as build:
            self.assertEqual(self.client.post("/api/v1/recommend/export", json=self.payload).status_code, 401)
            build.assert_not_called()

    def test_no_graph_ai_or_storage_calls(self):
        with patch("app.main.neo4j_client.run_rows") as graph, patch("app.main.deepseek_service.chat") as ai, patch("app.main.storage.save_strategy") as save:
            _, workbook = self.export()
            workbook.close()
            graph.assert_not_called()
            ai.assert_not_called()
            save.assert_not_called()

    def test_legacy_metadata_does_not_invent_generation_conditions(self):
        self.payload["meta"] = {"reference_year": 2023}
        _, workbook = self.export()
        self.addCleanup(workbook.close)
        notes = dict(workbook["方案说明"].values)
        self.assertIn("未使用当前表单补写", notes["生成条件"])
        self.assertIn("旧版", notes["版本提示"])
        self.assertNotIn("生成时位次", notes)

    def test_all_scope_preserves_order_and_ignores_page_filter_description(self):
        self.payload.update(scope="all", total_count=2)
        self.payload["items"].append({"institution": "第二所", "major": "其他专业"})
        _, workbook = self.export()
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["推荐清单"]["C3"].value, "第二所")
        self.assertNotIn("页面筛选", dict(workbook["方案说明"].values))

    def test_unrestricted_eligible_and_ineligible_remain_distinct(self):
        for state, expected in [("eligible", "不限选科"), ("unknown", "选科待核实"), ("ineligible", "选科不符")]:
            with self.subTest(state=state):
                self.payload["items"][0].update(subject_requirement_kind="unrestricted", subject_eligibility=state)
                _, workbook = self.export()
                self.assertEqual(workbook["推荐清单"]["N2"].value, expected)
                workbook.close()

    def test_invalid_requests_fail_before_workbook_build(self):
        invalid = [
            {"items": []}, {"items": self.payload["items"] * 301},
            {"total_count": 301}, {"scope": "all"}, {"scope": "pending"},
            {"filters": {"keyword": "x" * 2001}},
            {"items": [{"institution": "", "major": "专业"}]},
            {"items": [{"institution": "学校", "major": "专业", "plan_count": -1}]},
        ]
        with patch("app.main.build_recommendation_workbook") as build:
            for change in invalid:
                with self.subTest(change=str(change)[:80]):
                    self.assertEqual(self.client.post("/api/v1/recommend/export", json={**self.payload, **change}).status_code, 422)
            build.assert_not_called()

    def test_maximum_supported_export(self):
        self.payload.update(scope="all", total_count=300)
        self.payload["items"] *= 300
        _, workbook = self.export()
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["推荐清单"].max_row, 301)

    def test_generated_snapshot_is_independent_of_later_input_changes(self):
        client = Mock()
        client.run_rows.return_value = []
        request = RecommendationRequest(rank=11100, provinces=["上海"], ratio_chong=0)
        result = RecommenderService(client).recommend(request)
        expected = deepcopy(request.model_dump())
        request.rank = 99999
        request.provinces.append("北京")
        self.assertEqual(result["meta"]["input_snapshot"], expected)
        self.assertIsNotNone(datetime.fromisoformat(result["meta"]["generated_at"]).tzinfo)


if __name__ == "__main__":
    unittest.main()
