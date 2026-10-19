"""Raw-workbook readers and conservative normalization for the Zhejiang dataset."""
from __future__ import annotations

import hashlib
import math
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

DATASET_ID = "zhejiang-admissions-2019-2024-v2"
MISSING = {"", "缺失", "暂无", "无", "--", "-", "—", "nan", "None", "/"}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return unicodedata.normalize("NFKC", str(value)).replace("_x000D_", "").replace("\u200b", "").strip()


def name(value: Any) -> str:
    return re.sub(r"\s+", "", text(value))


def number(value: Any, integer: bool = False, positive: bool = False) -> int | float | None:
    value = text(value).replace(",", "")
    if value in MISSING:
        return None
    try:
        result = float(value)
        if not math.isfinite(result) or result < 0 or (positive and result == 0):
            return None
        if integer:
            return int(result) if result.is_integer() else None
        return int(result) if result.is_integer() else result
    except (ValueError, OverflowError):
        return None


def key(*parts: Any) -> str:
    return hashlib.sha256("\x1f".join(map(str, parts)).encode("utf-8")).hexdigest()


def catalog_code(raw: Any, description: Any) -> tuple[str, bool]:
    code = text(raw).upper()
    if re.fullmatch(r"\d{6}[TK]{0,2}", code):
        return code, False
    recovered = re.search(r"专业代码[:：]\s*(\d{6}[TKtk]{0,2})(?![0-9A-Za-z])", text(description))
    return (recovered[1].upper(), True) if recovered else (code, False)


def excel_records(path: Path, limit: int = 0) -> Iterator[tuple[str, int, dict[str, Any]]]:
    """Stream rows; detect the leading blank row in the subject-evaluation workbook."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book:
            header = None
            emitted = 0
            for index, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if header is None:
                    if not any(value is not None and text(value) for value in values):
                        continue
                    header = [text(value) for value in values]
                    continue
                if limit and emitted >= limit:
                    break
                if not any(value is not None and text(value) for value in values):
                    continue
                emitted += 1
                yield sheet.title, index, {h: v for h, v in zip(header, values) if h}
    finally:
        book.close()


SUBJECTS = {"物理": "物理", "化学": "化学", "生物": "生物", "政治": "政治", "思想政治": "政治", "历史": "历史", "地理": "地理", "技术": "技术", "物": "物理", "化": "化学", "生": "生物", "政": "政治", "史": "历史", "地": "地理", "技": "技术"}


UNRESTRICTED_SUBJECTS = {
    "", "不限", "无", "无要求", "没有要求", "无科目要求", "无选科要求", "无选考科目要求",
    "选科不限", "选考科目不限", "不限选考科目", "不提科目要求", "不提选考科目要求",
    "不作选科要求", "不作要求", "任意科目", "任何选科", "任何选考科目",
}


def subject_options(raw: str | None) -> list[list[str]] | None:
    """Blank source requirements mean unrestricted; absent source data is unknown."""
    if raw is None:
        return None
    value = name(raw).replace("思想政治", "政治")
    if value in {"未注明", "待确认", "待核实", "未知"}:
        return None
    if value in UNRESTRICTED_SUBJECTS:
        return [[]]
    value = re.sub(r"[()（）\s]", "", value)
    groups = re.split(r"或|/|\|", value)
    result = []
    for group in groups:
        tokens = [x for x in re.split(r"&|和|及|、|,|\+", group) if x]
        if not tokens or any(token not in SUBJECTS for token in tokens):
            return None
        result.append(sorted({SUBJECTS[token] for token in tokens}))
    return result


def subject_requirement_info(raw: str | None, *, plan_matched: bool = False) -> dict[str, Any]:
    # A matched plan's empty cell follows the user's no-requirement rule.
    # No matched plan provides no evidence about requirements.
    effective = "" if raw is None and plan_matched else raw
    options = subject_options(effective)
    kind = "unknown" if options is None else "unrestricted" if [] in options else "restricted"
    display = "不限选科" if kind == "unrestricted" else text(raw) or "待核实（未匹配到选科要求）"
    return {"effective": effective, "options": options, "kind": kind, "display": display}


def selected_subjects(raw: str) -> set[str]:
    raw = text(raw).strip("()").replace("思想政治", "政治")
    tokens = [x for x in re.split(r"[、,，/&+\s]+", raw) if x]
    if any(token not in SUBJECTS for token in tokens):
        raise ValueError("选考科目请使用物理、化学、生物、政治、历史、地理、技术，以逗号或顿号分隔")
    result = {SUBJECTS[token] for token in tokens}
    if len(result) > 3:
        raise ValueError("请填写最多三门选考科目；尚未填全时将展示待核实的探索候选")
    return result


def eligibility(requirement: str | None, selected: set[str], *, complete: bool = True) -> str:
    options = subject_options(requirement)
    if options is None:
        return "unknown"
    if any(set(group) <= selected for group in options):
        return "eligible"
    return "ineligible" if complete and selected else "unknown"
