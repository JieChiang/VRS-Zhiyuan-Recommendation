"""Export a bounded browser result snapshot without querying the graph or AI."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
import re
from typing import Annotated, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, FiniteFloat, StringConstraints, model_validator

from app.models import RecommendationRequest
from app.services.preference_filters import RECOMMENDATION_VERSION


Text = Annotated[str, StringConstraints(max_length=2000)]
PositiveInt = Annotated[int, Field(ge=1)]
CHINA_TIME = timezone(timedelta(hours=8))
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportItem(BaseModel):
    institution: Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=300)]
    major: Annotated[str, StringConstraints(strip_whitespace=True, min_length=1, max_length=1000)]
    year: int | None = Field(default=None, ge=2010, le=2100)
    layer: Literal["冲", "稳", "保"] | None = None
    min_rank: PositiveInt | None = None
    min_score: FiniteFloat | None = None
    plan_count: int | None = Field(default=None, ge=0)
    destination_province: Text | None = None
    destination_city: Text | None = None
    level: Text | None = None
    school_nature: Text | None = None
    subject_requirement_display: Text | None = None
    subject_requirement: Text | None = None
    subject_requirement_kind: Text | None = None
    subject_eligibility: Literal["eligible", "unknown", "ineligible"] | None = None
    subject_missing: list[Text] = Field(default_factory=list, max_length=7)
    source_file: Text | None = None
    source_sheet: Text | None = None
    source_row: PositiveInt | None = None
    plan_source: Text | None = None
    plan_source_row: PositiveInt | None = None
    fact_id: Text | None = None
    plan_fact_id: Text | None = None


class ExportMetadata(BaseModel):
    reference_year: int | None = None
    dataset_id: Text = ""
    recommendation_version: Text = ""
    generated_at: Text = ""
    input_snapshot: RecommendationRequest | None = None
    selected_subjects: list[Text] = Field(default_factory=list, max_length=7)
    warnings: list[Text] = Field(default_factory=list, max_length=100)
    qualification_note: Text = ""


class ExportFilters(BaseModel):
    keyword: Text = ""
    layer: Literal["", "冲", "稳", "保"] = ""
    level: Literal["", "college", "bachelor"] = ""
    nature: Literal["", "public", "private"] = ""


class RecommendationExportRequest(BaseModel):
    items: list[ExportItem] = Field(min_length=1, max_length=300)
    meta: ExportMetadata = Field(default_factory=ExportMetadata)
    scope: Literal["visible", "all"] = "visible"
    total_count: int = Field(ge=1, le=300)
    filters: ExportFilters = Field(default_factory=ExportFilters)

    @model_validator(mode="after")
    def check_counts(self):
        if len(self.items) > self.total_count or (self.scope == "all" and len(self.items) != self.total_count):
            raise ValueError("导出条数与推荐总数不一致，请刷新后重试")
        # Limit even optional input snapshots before writing workbook cells.
        if self.meta.input_snapshot and len(self.meta.input_snapshot.model_dump_json()) > 20000:
            raise ValueError("生成条件过长，请减少偏好内容")
        return self


def _text(value: str) -> str:
    # XML 1.0 cannot contain control characters or unpaired surrogates.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)[:32767]


def _append(sheet, values):
    sheet.append([_text(value) if isinstance(value, str) else value for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # Keep all user/source text literal, including leading '=' and error codes.
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _style(sheet, widths):
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1E40AF")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def build_recommendation_workbook(payload: RecommendationExportRequest) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "方案说明"
    meta = payload.meta
    counts = Counter(item.layer or "未分层" for item in payload.items)
    notes = [
        ("项目", "内容"),
        ("文件用途", "浙江历史志愿推荐参考，非正式志愿填报表。清单顺序为当前推荐展示顺序，不代表最终填报顺序。"),
        ("导出时间（北京时间）", datetime.now(CHINA_TIME).isoformat(timespec="seconds")),
        ("方案生成时间", meta.generated_at or "旧版结果未记录"),
        ("导出范围", "当前筛选结果" if payload.scope == "visible" else "全部推荐结果"),
        ("导出条数 / 推荐总数", f"{len(payload.items)} / {payload.total_count}"),
        ("导出分层统计", "；".join(f"{layer}：{counts[layer]}" for layer in ["冲", "稳", "保", "未分层"])),
        ("历史参考年份", meta.reference_year),
        ("数据集", meta.dataset_id or "未记录"),
        ("推荐规则版本", meta.recommendation_version or "未记录"),
        ("资料说明", "本文件保存浏览器中的推荐快照，未重新查询数据库。缺失数值留空，0 表示原值为 0。单独列出的待核实记录、目标院校补充列表和策略优化结果不包含在本清单中。"),
        ("使用说明", "冲稳保分层依据历史位次，不是录取保证；未导出未经校准的录取概率。历史选科和计划不等于当年报考资格，请核对当年官方招生章程。"),
    ]
    if payload.scope == "visible":
        filters = payload.filters
        notes.append(("页面筛选", "；".join([
            f"关键词：{filters.keyword or '不限'}", f"分层：{filters.layer or '全部'}",
            f"层次：{ {'college': '专科', 'bachelor': '本科'}.get(filters.level, '全部')}",
            f"性质：{ {'public': '公办', 'private': '民办'}.get(filters.nature, '全部')}",
        ])))
    request = meta.input_snapshot
    if request:
        notes.extend([
            ("生成时位次", request.rank), ("生成时分数", request.score),
            ("生成时选科", request.subject_combo or "未填写"),
            ("专业偏好", "、".join(request.majors) or "不限"),
            ("意向省市", "、".join(request.provinces) or "不限"),
            ("目标院校", "、".join(request.target_institutions) or "未指定"),
            ("目标条数", request.target_count),
            ("冲 / 稳 / 保比例", f"{request.ratio_chong} / {request.ratio_wen} / {request.ratio_bao}"),
            ("候选检索范围", "全部历史位次" if request.candidate_scope == "all" else f"位次上下 {request.rank_window}"),
        ])
    else:
        notes.extend([
            ("生成条件", "旧版结果未保存完整生成条件；未使用当前表单补写。可重新生成推荐后导出。"),
            ("结果中记录的选科", "、".join(meta.selected_subjects) or "未记录"),
        ])
    if meta.recommendation_version != RECOMMENDATION_VERSION:
        notes.append(("版本提示", "本方案使用旧版或未知推荐规则，建议重新生成推荐。"))
    if meta.qualification_note:
        notes.append(("选科核验说明", meta.qualification_note))
    notes.extend(("生成提示", warning) for warning in meta.warnings)
    for row in notes:
        _append(summary, row)
    _style(summary, [27, 110])
    for row in range(2, summary.max_row + 1):
        summary.row_dimensions[row].height = 48

    sheet = workbook.create_sheet("推荐清单")
    _append(sheet, ["序号", "层次", "院校", "招生专业", "参考年份", "历史最低分", "历史最低位次", "历史计划数",
                    "院校省份", "院校城市", "办学层次", "办学性质", "选科要求", "选科核验", "待确认科目",
                    "录取来源文件", "录取来源工作表", "录取来源行号", "计划来源文件", "计划来源行号", "录取事实ID", "计划事实ID"])
    for index, item in enumerate(payload.items, 1):
        status = {"eligible": "选科符合计划原文", "ineligible": "选科不符"}.get(item.subject_eligibility, "选科待核实")
        if item.subject_requirement_kind == "unrestricted" and item.subject_eligibility == "eligible":
            status = "不限选科"
        _append(sheet, [index, item.layer, item.institution, item.major, item.year, item.min_score, item.min_rank, item.plan_count,
                        item.destination_province, item.destination_city, item.level, item.school_nature,
                        item.subject_requirement_display or item.subject_requirement or "待核实", status, "、".join(item.subject_missing),
                        item.source_file, item.source_sheet, item.source_row, item.plan_source, item.plan_source_row,
                        item.fact_id, item.plan_fact_id])
        sheet.row_dimensions[index + 1].height = 45
        color = {"冲": "FFF1F2", "稳": "EFF6FF", "保": "F0FDF4"}.get(item.layer, "F8FAFC")
        for cell in sheet[index + 1]:
            cell.fill = PatternFill("solid", fgColor=color)
    _style(sheet, [8, 9, 28, 42, 12, 14, 16, 14, 14, 14, 14, 14, 28, 25, 20, 30, 24, 16, 30, 16, 32, 32])
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
