"""Offline Chromium UI/download smoke test using optional Playwright tooling.

Only serves test pages, local assets and the export endpoint. No real graph or AI
requests, persistent browser profile or changes to the user's browser data.
"""
from __future__ import annotations

import json
import os
from pathlib import Path
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "storage" / "export-verification"
OUTPUT.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(ROOT / "storage" / "browser-runtime"))
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from app.main import app


def main():
    client = TestClient(app)
    client.cookies.set("vrs_auth", "1")

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_GET(self):
            if self.path == "/api/v1/setup/status":
                component = {"name": "离线测试", "title": "测试就绪", "state": "ok", "ready": True,
                             "configured": True, "available": True, "verified": True, "detail": "合成测试数据", "docs_url": "#"}
                body = json.dumps({"data_ready": True, "components": {key: component for key in ("neo4j", "deepseek", "ocr")}}).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(body)
            elif self.path == "/recommend" or self.path.startswith("/static/"):
                self.reply(client.get(self.path))
            else:
                self.send_error(404)

        def do_POST(self):
            if self.path != "/api/v1/recommend/export":
                self.send_error(404)
                return
            body = self.rfile.read(int(self.headers.get("Content-Length", "0")))
            self.reply(client.post(self.path, content=body, headers={"Content-Type": "application/json"}))

        def reply(self, response):
            self.send_response(response.status_code)
            for name in ("content-type", "content-disposition", "cache-control"):
                if name in response.headers:
                    self.send_header(name, response.headers[name])
            self.end_headers()
            self.wfile.write(response.content)

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            context = browser.new_context(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
            page = context.new_page()
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            base = f"http://127.0.0.1:{server.server_port}"
            context.route("**/*", lambda route: route.continue_() if route.request.url.startswith(base + "/") else route.abort())
            page.goto(base + "/recommend")
            with page.expect_download() as download:
                page.evaluate((ROOT / "tests" / "recommendation_export_browser.js").read_text(encoding="utf-8"))
                result = page.evaluate("window.exportVerification")
                assert not result.get("error"), result
            download.value.save_as(OUTPUT / "recommendations.xlsx")
            workbook = load_workbook(OUTPUT / "recommendations.xlsx")
            assert workbook["推荐清单"].max_row == 5
            assert workbook["推荐清单"]["H2"].value == 0
            assert workbook["推荐清单"]["H3"].value is None
            assert dict(workbook["方案说明"].values)["生成时位次"] == 11100
            workbook.close()
            page.screenshot(path=str(OUTPUT / "desktop.png"), full_page=True)
            page.set_viewport_size({"width": 390, "height": 844})
            page.wait_for_function("""getComputedStyle(document.querySelector('.app-main')).marginLeft === '0px'
              && document.getElementById('sidebar').getBoundingClientRect().right <= 1""")
            page.locator("#mobileNavToggle").click()
            page.wait_for_function("document.getElementById('sidebar').getBoundingClientRect().left >= 0")
            page.keyboard.press("Escape")
            page.wait_for_function("document.getElementById('sidebar').getBoundingClientRect().right <= 1")
            page.screenshot(path=str(OUTPUT / "mobile.png"), full_page=True)
            overflow = page.evaluate("""Array.from(document.querySelectorAll('#mainContent *')).filter(el => {
              const rect = el.getBoundingClientRect(); return rect.width && rect.right > innerWidth + 1;
            }).map(el => ({tag:el.tagName, id:el.id, className:el.className, width:el.getBoundingClientRect().width})).slice(0, 12)""")
            assert page.evaluate("document.documentElement.scrollWidth") <= 390, f"Mobile page overflows horizontally: {overflow}"
            assert page.locator("#exportRecommendationsBtn").bounding_box()["width"] >= 44
            assert not errors, errors
            context.close()
            browser.close()
            print(json.dumps({"browser_checks": result["checks"], "download": "passed", "mobile_overflow": False,
                              "output": str(OUTPUT)}, ensure_ascii=False))
    finally:
        server.shutdown()
        server.server_close()
        client.close()


if __name__ == "__main__":
    main()
