"""Exercise the real Qt shell without visible windows or external services."""
from __future__ import annotations

import json
import os
from pathlib import Path
import socket
import sys
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "storage" / "desktop-verification"
OUTPUT.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(ROOT))
# Use the native Windows platform and hide only the test windows, not Chromium.
os.environ.setdefault("QT_QPA_PLATFORM", "windows" if sys.platform == "win32" else "offscreen")

from PyQt6.QtCore import QCoreApplication, QEvent, QEventLoop, QTimer
from openpyxl import load_workbook

from app.desktop_runtime import DesktopServer, is_vrs
from app.desktop_window import BrowserManager, create_application
from app.main import configuration_checks


def main():
    app = create_application()
    app.setQuitOnLastWindowClosed(False)
    with socket.socket() as socket_probe:
        socket_probe.bind(("127.0.0.1", 0))
        port = socket_probe.getsockname()[1]
    server = DesktopServer(port)
    component = {"name": "离线测试", "title": "就绪", "state": "ok", "ready": True,
                 "configured": True, "available": True, "verified": True, "detail": "测试", "docs_url": "#"}
    status = {"data_ready": True, "components": {key: component for key in ("neo4j", "deepseek", "ocr")}}
    downloads, external, crashes = [], [], []
    manager = None

    def tick():
        loop = QEventLoop()
        QTimer.singleShot(25, loop.quit)
        loop.exec()

    def until(predicate, timeout=20):
        deadline = time.monotonic() + timeout
        while not predicate():
            if crashes:
                raise AssertionError(f"Renderer crashed: {crashes}")
            if time.monotonic() > deadline:
                raise AssertionError("Qt verification timed out")
            tick()

    def js(window, code):
        result = []
        window.web.page().runJavaScript(code, result.append)
        until(lambda: bool(result))
        return result[0]

    def choose(suggested):
        downloads.append(suggested.name)
        return str(OUTPUT / "recommendations.xlsx")

    try:
        with patch.object(configuration_checks, "status", return_value=status):
            server.start()
            reused = DesktopServer(port)
            reused.start()
            assert reused.reused
            reused.stop()
            assert is_vrs(server.url), "Closing reused server must not stop existing service"
            manager = BrowserManager(app, server.url, OUTPUT / "browser", hidden=True,
                                     choose_download=choose, external_opener=external.append)
            window = manager.open_window(server.url + "/login")
            window.web.page().renderProcessTerminated.connect(lambda *args: crashes.append(str(args)))
            until(lambda: js(window, "Boolean(document.getElementById('username'))"))
            js(window, """localStorage.clear(); fetch('/login', {method:'POST',
              headers:{'Content-Type':'application/x-www-form-urlencoded'},
              body:'username=desktop-test&password=temporary-test'}).then(() => location.assign('/recommend'));""")
            until(lambda: js(window, "location.pathname === '/recommend' && document.readyState === 'complete' && typeof render === 'function'"))
            js(window, (ROOT / "tests" / "recommendation_export_browser.js").read_text(encoding="utf-8"))
            until(lambda: js(window, "Boolean(window.exportVerification?.done)"))
            result = json.loads(js(window, "JSON.stringify(window.exportVerification)"))
            assert not result.get("error"), result
            until(lambda: downloads and not manager.downloads)
            workbook = load_workbook(OUTPUT / "recommendations.xlsx")
            assert workbook["推荐清单"].max_row == 5
            assert dict(workbook["方案说明"].values)["生成时位次"] == 11100
            workbook.close()
            # Persistence after navigation, and shared cookies/storage in internal popups.
            window.web.reload()
            until(lambda: js(window, "typeof recommendationExportPayload === 'function' && recommendationExportPayload().items.length === 4"))
            js(window, "window.open('/compare', 'vrs-compare');")
            until(lambda: "/compare" in manager.windows)
            compare = manager.windows["/compare"]
            until(lambda: js(compare, "location.pathname === '/compare' && document.readyState === 'complete'"))
            assert js(compare, "JSON.parse(localStorage.getItem('vrs.recommendResult')).all.length") == 4
            # The login cookie is HttpOnly; protected page access proves sharing.
            assert js(compare, "document.title.includes('对比')")
            assert manager.open_window(server.url + "/compare") is compare
            js(window, "window.open('https://example.org/desktop-test', '_blank');")
            until(lambda: bool(external))
            assert external == ["https://example.org/desktop-test"]
            manager.open_external("file:///C:/should-not-open")
            assert len(external) == 1
            compare.close()
            until(lambda: "/compare" not in manager.windows)
            # Cancel a save dialog without opening a native modal during the test.
            cancelled = []
            manager.choose_download = lambda _: cancelled.append(True) or ""
            js(window, "exportRecommendations();")
            until(lambda: bool(cancelled))
            assert not manager.downloads
            until(lambda: js(window, "!document.getElementById('exportRecommendationsBtn').disabled"))
            js(window, "document.getElementById('toastContainer').replaceChildren()")
            tick()
            window.grab().save(str(OUTPUT / "desktop.png"))
            assert not crashes
            print(json.dumps({"qt_browser_checks": result["checks"], "login": "passed", "download": "passed",
                              "reload_and_popup_state": "passed", "cancel_download": "passed", "renderer_crashes": 0}, ensure_ascii=False))
    finally:
        if manager:
            manager.dispose()
        server.stop()
    assert not is_vrs(server.url), "Owned server must stop when desktop closes"
    print("Owned server stopped; existing server reuse preserved.")


def verify_entry():
    """Run the real entry point and close its hidden window after first load."""
    from desktop_app import main as desktop_main

    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    loaded = []

    class HiddenManager(BrowserManager):
        def __init__(self, app, url, storage):
            super().__init__(app, url, OUTPUT / "entry-browser", hidden=True)
            QTimer.singleShot(15000, lambda: app.exit(2))

        def open_window(self, url):
            window = super().open_window(url)

            def ready(ok):
                loaded.append(ok)
                QTimer.singleShot(100, self.close)

            window.web.loadFinished.connect(ready)
            return window

    with patch("app.desktop_window.BrowserManager", HiddenManager):
        code = desktop_main(["--port", str(port)])
    assert code == 0 and loaded == [True], (code, loaded)
    assert not is_vrs(f"http://127.0.0.1:{port}")
    print("Real desktop entry: loaded login, closed window, stopped owned backend, exited 0.")


if __name__ == "__main__":
    verify_entry() if "--entry" in sys.argv else main()
